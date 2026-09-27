"""Add SKUs from an Excel file to the sheet's SKUs tab (SKUs already there are skipped).

    python import_skus.py "D:\\Smartivity_SKU_Price_Tracking_Input.xlsx"            # preview only
    python import_skus.py "D:\\Smartivity_SKU_Price_Tracking_Input.xlsx" --upload   # add to the sheet

Needs SHEET_WEBAPP_URL and SHEET_TOKEN for --upload. Requires: pip install openpyxl
"""
import re
import sys

from openpyxl import load_workbook

from tracker import sheet

# Excel column header -> SKUs tab column header (same name unless listed)
RENAME = {}
SKIP_SKUS = {"SMRT1109A"}          # duplicate of SMRT1109


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_rows(path):
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [clean(h) for h in next(it)]
    rows = []
    for values in it:
        r = {RENAME.get(h, h): clean(v) for h, v in zip(head, values) if h}
        r["SKU"] = re.sub(r"\s+", "", r.get("SKU", ""))
        if not re.fullmatch(r"SMRT\d+[A-Z]?", r["SKU"]) or r["SKU"] in SKIP_SKUS:
            continue
        r["Track (Y/N)"] = "Y"
        rows.append(r)
    return rows


if __name__ == "__main__":
    rows = read_rows(sys.argv[1])
    print(f"{len(rows)} SKUs read; with ASIN: {sum(bool(r.get('Amazon ASIN(s)')) for r in rows)}")
    if "--upload" in sys.argv:
        print(sheet._post({"action": "add_skus", "rows": rows}, timeout=120)["message"])
